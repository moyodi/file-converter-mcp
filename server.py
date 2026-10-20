"""
file-converter-mcp
업로드 없이, 유료 결제 없이, 대화 한 줄로 끝나는 파일 변환 MCP 서버

1단계: 이미지 변환 (PNG · JPG · WebP · HEIC · GIF · BMP · TIFF)
2단계: 영상·음성 변환 (MOV · MP4 · WebM · GIF · MP3 · M4A · WAV)
3단계: 문서 변환 (HWP · DOCX · PPTX · XLSX → PDF, HWP → DOCX, PDF → DOCX·이미지, 이미지 → PDF, XLSX ↔ CSV)
- 파일은 내 컴퓨터 밖으로 나가지 않습니다.
- 원본은 절대 덮어쓰지 않고, 변환 결과는 별도 폴더에 저장합니다.
"""

from __future__ import annotations

import csv
import os
import platform
import shutil
import subprocess
import tempfile
import urllib.request
import sys
from pathlib import Path

from mcp.server.mcpserver import MCPServer
from mcp.server.mcpserver.exceptions import ToolError
from PIL import Image, ImageOps

# 아이폰 사진(HEIC) 지원 — 설치돼 있으면 자동으로 켜짐
try:
    from pillow_heif import register_heif_opener

    register_heif_opener()
    HEIC_SUPPORTED = True
except ImportError:
    HEIC_SUPPORTED = False


mcp = MCPServer(
    name="file-convert",
    instructions=(
        "로컬 파일 변환 도구입니다. 사용자가 이미지·영상·음성의 확장자 변경, "
        "용량 줄이기, 크기 줄이기, 폴더 일괄 변환, 영상에서 소리 추출, "
        "영상을 GIF로 만들기를 요청하면 이 도구를 사용하세요. "
        "이미지는 convert_file / convert_folder, 영상·음성은 convert_video / convert_video_folder, "
        "문서(HWP·DOCX·PPTX·XLSX·PDF·CSV)는 convert_document / convert_document_folder, "
        "여러 이미지를 PDF 하나로 묶을 때는 images_to_pdf를 쓰세요. "
        "경로는 사용자 컴퓨터의 절대 경로(~ 사용 가능)를 넘기세요."
    ),
)

# 바꿀 수 있는 형식: 사용자가 쓰는 이름 → Pillow 저장 형식
OUTPUT_FORMATS = {
    "jpg": "JPEG",
    "jpeg": "JPEG",
    "png": "PNG",
    "webp": "WEBP",
    "gif": "GIF",
    "bmp": "BMP",
    "tiff": "TIFF",
}

# 읽을 수 있는 확장자
INPUT_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
if HEIC_SUPPORTED:
    INPUT_EXTENSIONS |= {".heic", ".heif"}

# 투명 배경을 지원하지 않는 형식
NO_ALPHA = {"JPEG", "BMP"}


def _log(msg: str) -> None:
    # stdio MCP는 표준출력을 통신에 쓰므로, 로그는 반드시 stderr로
    print(msg, file=sys.stderr)


def _human_size(num_bytes: int) -> str:
    size = float(num_bytes)
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            return f"{size:.1f}{unit}" if unit != "B" else f"{int(size)}B"
        size /= 1024
    return f"{size:.1f}GB"


def _unique_path(path: Path) -> Path:
    """같은 이름 파일이 있으면 _1, _2 를 붙여서 덮어쓰기 방지"""
    if not path.exists():
        return path
    i = 1
    while True:
        candidate = path.with_name(f"{path.stem}_{i}{path.suffix}")
        if not candidate.exists():
            return candidate
        i += 1


def _normalize_format(to_format: str) -> tuple[str, str]:
    key = to_format.lower().lstrip(".").strip()
    if key not in OUTPUT_FORMATS:
        raise ToolError(
            f"'{to_format}' 형식은 아직 지원하지 않아요. "
            f"가능한 형식: {', '.join(sorted(set(OUTPUT_FORMATS)))}"
        )
    ext = "jpg" if key == "jpeg" else key
    return OUTPUT_FORMATS[key], ext


def _convert_one(
    src: Path,
    out_dir: Path,
    pil_format: str,
    ext: str,
    quality: int,
    max_width: int | None,
) -> dict:
    with Image.open(src) as img:
        # 휴대폰 사진이 옆으로 눕는 문제 방지 (EXIF 회전값 반영)
        img = ImageOps.exif_transpose(img)

        # 가로 크기 제한 (비율 유지, 확대는 하지 않음)
        if max_width and img.width > max_width:
            ratio = max_width / img.width
            img = img.resize((max_width, round(img.height * ratio)), Image.LANCZOS)

        # JPG처럼 투명도가 없는 형식은 흰 배경으로 합성
        if pil_format in NO_ALPHA and img.mode in ("RGBA", "LA", "P"):
            img = img.convert("RGBA")
            background = Image.new("RGB", img.size, (255, 255, 255))
            background.paste(img, mask=img.split()[-1])
            img = background
        elif pil_format in NO_ALPHA and img.mode != "RGB":
            img = img.convert("RGB")

        save_kwargs: dict = {}
        if pil_format in ("JPEG", "WEBP"):
            save_kwargs["quality"] = quality
        if pil_format == "JPEG":
            save_kwargs["optimize"] = True
        if pil_format == "PNG":
            save_kwargs["optimize"] = True

        out_dir.mkdir(parents=True, exist_ok=True)
        dst = _unique_path(out_dir / f"{src.stem}.{ext}")
        img.save(dst, pil_format, **save_kwargs)
        width, height = img.size

    before, after = src.stat().st_size, dst.stat().st_size
    return {
        "원본": str(src),
        "결과": str(dst),
        "크기": f"{width}x{height}",
        "용량": f"{_human_size(before)} → {_human_size(after)}",
        "_before": before,
        "_after": after,
    }


@mcp.tool()
def list_formats() -> dict:
    """변환할 수 있는 이미지·영상·음성 형식 목록을 알려줍니다."""
    return {
        "이미지": {
            "읽기 가능": sorted(e.lstrip(".") for e in INPUT_EXTENSIONS),
            "변환 가능": sorted({"jpg", "png", "webp", "gif", "bmp", "tiff"}),
            "HEIC(아이폰 사진) 지원": HEIC_SUPPORTED,
        },
        "영상·음성": {
            "읽기 가능": sorted(e.lstrip(".") for e in MEDIA_INPUT_EXTENSIONS),
            "변환 가능": sorted(VIDEO_OUTPUTS),
            "엔진(FFmpeg) 준비됨": _find_ffmpeg() is not None,
        },
        "문서": {
            "변환 경로": DOC_ROUTES_HELP,
            "LibreOffice 설치됨 (HWP·오피스 → PDF에 필요)": _find_soffice() is not None,
        },
    }


@mcp.tool()
def convert_file(
    path: str,
    to_format: str,
    quality: int = 90,
    max_width: int | None = None,
    output_dir: str | None = None,
) -> dict:
    """
    이미지 파일 하나를 다른 형식으로 변환합니다. 원본은 그대로 둡니다.

    Args:
        path: 변환할 파일 경로 (예: ~/Desktop/photo.heic)
        to_format: 바꿀 형식 (jpg, png, webp, gif, bmp, tiff)
        quality: 화질 1~100 (jpg·webp에만 적용, 기본 90)
        max_width: 가로 최대 픽셀. 지정하면 비율 유지하며 줄임 (예: 1080)
        output_dir: 저장할 폴더. 비우면 원본 옆 'converted_형식' 폴더에 저장
    """
    src = Path(path).expanduser().resolve()
    if not src.is_file():
        raise ToolError(f"파일을 찾을 수 없어요: {src}")
    if src.suffix.lower() not in INPUT_EXTENSIONS:
        raise ToolError(f"이미지 파일이 아니거나 아직 지원하지 않는 형식이에요: {src.name}")

    pil_format, ext = _normalize_format(to_format)
    quality = max(1, min(100, quality))
    out_dir = Path(output_dir).expanduser().resolve() if output_dir else src.parent / f"converted_{ext}"

    result = _convert_one(src, out_dir, pil_format, ext, quality, max_width)
    _log(f"[file-convert] {src.name} → {result['결과']}")
    return {k: v for k, v in result.items() if not k.startswith("_")}


@mcp.tool()
def convert_folder(
    folder: str,
    to_format: str,
    quality: int = 90,
    max_width: int | None = None,
    include_subfolders: bool = False,
    output_dir: str | None = None,
) -> dict:
    """
    폴더 안의 이미지를 한 번에 변환합니다. 원본은 그대로 둡니다.

    Args:
        folder: 이미지가 들어 있는 폴더 경로 (예: ~/Desktop/촬영본)
        to_format: 바꿀 형식 (jpg, png, webp, gif, bmp, tiff)
        quality: 화질 1~100 (jpg·webp에만 적용, 기본 90)
        max_width: 가로 최대 픽셀. 지정하면 비율 유지하며 줄임 (예: 1080)
        include_subfolders: 하위 폴더까지 포함할지 여부
        output_dir: 저장할 폴더. 비우면 대상 폴더 안 'converted_형식' 폴더에 저장
    """
    base = Path(folder).expanduser().resolve()
    if not base.is_dir():
        raise ToolError(f"폴더를 찾을 수 없어요: {base}")

    pil_format, ext = _normalize_format(to_format)
    quality = max(1, min(100, quality))
    out_root = Path(output_dir).expanduser().resolve() if output_dir else base / f"converted_{ext}"

    pattern = "**/*" if include_subfolders else "*"
    files = sorted(
        p for p in base.glob(pattern)
        if p.is_file()
        and p.suffix.lower() in INPUT_EXTENSIONS
        and out_root not in p.parents
        # 이전에 만든 converted_ 폴더 속 결과물은 다시 변환하지 않음
        and not any(part.startswith("converted_") for part in p.relative_to(base).parts[:-1])
    )
    if not files:
        return {"메시지": "변환할 이미지가 없어요.", "폴더": str(base)}

    done, failed = [], []
    total_before = total_after = 0
    for src in files:
        # 하위 폴더 구조를 결과 폴더에도 그대로 유지
        target_dir = out_root / src.parent.relative_to(base)
        try:
            r = _convert_one(src, target_dir, pil_format, ext, quality, max_width)
            total_before += r["_before"]
            total_after += r["_after"]
            done.append(Path(r["결과"]).name)
        except Exception as e:  # 한 장이 깨져도 나머지는 계속
            failed.append({"파일": src.name, "이유": str(e)})

    saved = (1 - total_after / total_before) * 100 if total_before else 0
    _log(f"[file-convert] {len(done)}/{len(files)}개 변환 → {out_root}")
    return {
        "변환 완료": f"{len(done)}개",
        "실패": failed,
        "저장 위치": str(out_root),
        "전체 용량": f"{_human_size(total_before)} → {_human_size(total_after)} ({saved:.0f}% 절감)"
        if saved >= 0
        else f"{_human_size(total_before)} → {_human_size(total_after)} ({-saved:.0f}% 증가)",
        "파일 목록": done[:50],
    }


# ─────────────────────────────────────────────
# 2단계: 영상·음성 변환 (FFmpeg)
# ─────────────────────────────────────────────

MEDIA_INPUT_EXTENSIONS = {
    # 영상
    ".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm", ".wmv", ".flv",
    ".mpg", ".mpeg", ".3gp", ".gif",
    # 음성
    ".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg", ".opus",
}
AUDIO_ONLY_INPUTS = {".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg", ".opus"}

VIDEO_OUTPUTS = {"mp4", "mov", "webm", "gif", "mp3", "m4a", "wav"}
AUDIO_OUTPUTS = {"mp3", "m4a", "wav"}

# 화질 단계 → FFmpeg CRF 값 (낮을수록 고화질·대용량)
QUALITY_CRF = {
    "high": {"h264": 20, "vp9": 28},
    "medium": {"h264": 26, "vp9": 34},
    "low": {"h264": 32, "vp9": 40},
}
QUALITY_ALIASES = {"고화질": "high", "보통": "medium", "중간": "medium", "저화질": "low"}

GIF_DEFAULT_WIDTH = 480
GIF_DEFAULT_SECONDS = 10


def _find_ffmpeg() -> str | None:
    """pip로 설치한 FFmpeg를 먼저 쓰고, 없으면 컴퓨터에 깔린 FFmpeg를 찾음"""
    try:
        import imageio_ffmpeg

        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        return shutil.which("ffmpeg")


def _has_audio(ffmpeg: str, src: Path) -> bool:
    """파일에 소리 트랙이 있는지 확인"""
    info = subprocess.run([ffmpeg, "-hide_banner", "-i", str(src)], capture_output=True, text=True).stderr
    return "Audio:" in info


def _scale_filter(max_width: int | None) -> str | None:
    # 원본보다 키우지 않고, 높이는 짝수로 맞춤(인코더 요구사항)
    if not max_width:
        return None
    return f"scale='min({max_width},iw)':-2"


def _build_ffmpeg_args(
    src: Path,
    dst: Path,
    fmt: str,
    quality: str,
    max_width: int | None,
    start: float | None,
    duration: float | None,
    fps: int | None,
    mute: bool,
) -> list[str]:
    args = ["-hide_banner", "-loglevel", "error", "-n"]
    if start:
        args += ["-ss", str(start)]
    args += ["-i", str(src)]
    if duration:
        args += ["-t", str(duration)]

    if fmt in AUDIO_OUTPUTS:
        args += ["-vn"]
        if fmt == "mp3":
            args += ["-c:a", "libmp3lame", "-q:a", {"high": "0", "medium": "2", "low": "5"}[quality]]
        elif fmt == "m4a":
            args += ["-c:a", "aac", "-b:a", {"high": "256k", "medium": "160k", "low": "96k"}[quality]]
        else:  # wav
            args += ["-c:a", "pcm_s16le"]
        return args + [str(dst)]

    if fmt == "gif":
        width = max_width or GIF_DEFAULT_WIDTH
        rate = fps or 12
        # 팔레트를 만들어 쓰면 GIF 색이 훨씬 깨끗해짐
        vf = (
            f"fps={rate},scale='min({width},iw)':-2:flags=lanczos,"
            "split[a][b];[a]palettegen[p];[b][p]paletteuse"
        )
        return args + ["-vf", vf, "-loop", "0", str(dst)]

    filters = [f for f in (_scale_filter(max_width), f"fps={fps}" if fps else None) if f]
    if filters:
        args += ["-vf", ",".join(filters)]

    if fmt == "webm":
        args += ["-c:v", "libvpx-vp9", "-crf", str(QUALITY_CRF[quality]["vp9"]), "-b:v", "0", "-row-mt", "1"]
        args += ["-an"] if mute else ["-c:a", "libopus", "-b:a", "128k"]
    else:  # mp4, mov → 어디서나 재생되는 H.264
        args += [
            "-c:v", "libx264", "-crf", str(QUALITY_CRF[quality]["h264"]),
            "-preset", "medium", "-pix_fmt", "yuv420p",
        ]
        args += ["-an"] if mute else ["-c:a", "aac", "-b:a", "160k"]
        args += ["-movflags", "+faststart"]  # 웹에서 바로 재생 시작
    return args + [str(dst)]


def _convert_media_one(
    src: Path,
    out_dir: Path,
    fmt: str,
    quality: str,
    max_width: int | None,
    start: float | None,
    duration: float | None,
    fps: int | None,
    mute: bool,
) -> dict:
    ffmpeg = _find_ffmpeg()
    if not ffmpeg:
        raise ToolError(
            "영상 엔진(FFmpeg)을 찾을 수 없어요. 터미널에서 "
            "`.venv/bin/pip install -r requirements.txt` 를 다시 실행해주세요."
        )
    if src.suffix.lower() in AUDIO_ONLY_INPUTS and fmt not in AUDIO_OUTPUTS:
        raise ToolError(f"음성 파일({src.name})은 mp3·m4a·wav로만 바꿀 수 있어요.")
    if fmt in AUDIO_OUTPUTS and not _has_audio(ffmpeg, src):
        raise ToolError(f"소리가 없는 영상이라 {fmt}로 추출할 수 없어요: {src.name}")

    note = None
    if fmt == "gif" and not duration:
        duration = GIF_DEFAULT_SECONDS
        note = f"GIF는 용량이 커서 앞부분 {GIF_DEFAULT_SECONDS}초만 만들었어요. 더 길게 원하면 duration을 지정하세요."

    out_dir.mkdir(parents=True, exist_ok=True)
    dst = _unique_path(out_dir / f"{src.stem}.{fmt}")
    args = _build_ffmpeg_args(src, dst, fmt, quality, max_width, start, duration, fps, mute)

    proc = subprocess.run([ffmpeg, *args], capture_output=True, text=True)
    if proc.returncode != 0 or not dst.exists():
        dst.unlink(missing_ok=True)
        err = (proc.stderr or "알 수 없는 오류").strip().splitlines()[-1][:300]
        raise ToolError(f"{src.name} 변환 실패: {err}")

    before, after = src.stat().st_size, dst.stat().st_size
    result = {
        "원본": str(src),
        "결과": str(dst),
        "용량": f"{_human_size(before)} → {_human_size(after)}",
        "_before": before,
        "_after": after,
    }
    if note:
        result["참고"] = note
    return result


def _normalize_media_args(to_format: str, quality: str) -> tuple[str, str]:
    fmt = to_format.lower().lstrip(".").strip()
    if fmt == "mpeg4":
        fmt = "mp4"
    if fmt not in VIDEO_OUTPUTS:
        raise ToolError(
            f"'{to_format}' 형식은 아직 지원하지 않아요. 가능한 형식: {', '.join(sorted(VIDEO_OUTPUTS))}"
        )
    q = QUALITY_ALIASES.get(quality, quality).lower()
    if q not in QUALITY_CRF:
        raise ToolError("quality는 high(고화질), medium(보통), low(저화질) 중 하나로 지정해주세요.")
    return fmt, q


@mcp.tool()
def convert_video(
    path: str,
    to_format: str,
    quality: str = "medium",
    max_width: int | None = None,
    start: float | None = None,
    duration: float | None = None,
    fps: int | None = None,
    mute: bool = False,
    output_dir: str | None = None,
) -> dict:
    """
    영상·음성 파일 하나를 변환합니다. 원본은 그대로 둡니다.
    MOV→MP4, 영상→GIF(움짤), 영상→MP3(소리 추출), 영상 용량 줄이기 등에 사용하세요.

    Args:
        path: 변환할 파일 경로 (예: ~/Desktop/IMG_0001.MOV)
        to_format: mp4, mov, webm, gif, mp3, m4a, wav
        quality: high(고화질) / medium(보통, 기본) / low(저화질·작은 용량)
        max_width: 가로 최대 픽셀. 지정하면 비율 유지하며 줄임 (예: 1280). GIF 기본값 480
        start: 시작 지점(초). 예: 5 → 5초부터
        duration: 길이(초). 예: 3 → 3초만. GIF는 지정 안 하면 앞 10초
        fps: 초당 프레임 수. GIF 기본값 12
        mute: True면 소리 없이 저장 (mp4·mov·webm)
        output_dir: 저장할 폴더. 비우면 원본 옆 'converted_형식' 폴더
    """
    src = Path(path).expanduser().resolve()
    if not src.is_file():
        raise ToolError(f"파일을 찾을 수 없어요: {src}")
    if src.suffix.lower() not in MEDIA_INPUT_EXTENSIONS:
        raise ToolError(f"영상·음성 파일이 아니거나 아직 지원하지 않는 형식이에요: {src.name}")

    fmt, q = _normalize_media_args(to_format, quality)
    out_dir = Path(output_dir).expanduser().resolve() if output_dir else src.parent / f"converted_{fmt}"

    result = _convert_media_one(src, out_dir, fmt, q, max_width, start, duration, fps, mute)
    _log(f"[file-convert] {src.name} → {result['결과']}")
    return {k: v for k, v in result.items() if not k.startswith("_")}


@mcp.tool()
def convert_video_folder(
    folder: str,
    to_format: str,
    quality: str = "medium",
    max_width: int | None = None,
    mute: bool = False,
    include_subfolders: bool = False,
    output_dir: str | None = None,
) -> dict:
    """
    폴더 안의 영상·음성 파일을 한 번에 변환합니다. 원본은 그대로 둡니다.

    Args:
        folder: 영상이 들어 있는 폴더 경로 (예: ~/Desktop/촬영본)
        to_format: mp4, mov, webm, gif, mp3, m4a, wav
        quality: high(고화질) / medium(보통, 기본) / low(저화질·작은 용량)
        max_width: 가로 최대 픽셀 (예: 1280)
        mute: True면 소리 없이 저장
        include_subfolders: 하위 폴더까지 포함할지 여부
        output_dir: 저장할 폴더. 비우면 대상 폴더 안 'converted_형식' 폴더
    """
    base = Path(folder).expanduser().resolve()
    if not base.is_dir():
        raise ToolError(f"폴더를 찾을 수 없어요: {base}")

    fmt, q = _normalize_media_args(to_format, quality)
    out_root = Path(output_dir).expanduser().resolve() if output_dir else base / f"converted_{fmt}"

    pattern = "**/*" if include_subfolders else "*"
    files = sorted(
        p for p in base.glob(pattern)
        if p.is_file()
        and p.suffix.lower() in MEDIA_INPUT_EXTENSIONS
        and p.suffix.lower() != ".gif"  # 폴더 일괄 변환에서는 GIF 이미지를 영상으로 보지 않음
        and out_root not in p.parents
        and not any(part.startswith("converted_") for part in p.relative_to(base).parts[:-1])
    )
    if not files:
        return {"메시지": "변환할 영상·음성 파일이 없어요.", "폴더": str(base)}

    done, failed = [], []
    total_before = total_after = 0
    for src in files:
        target_dir = out_root / src.parent.relative_to(base)
        try:
            r = _convert_media_one(src, target_dir, fmt, q, max_width, None, None, None, mute)
            total_before += r["_before"]
            total_after += r["_after"]
            done.append(Path(r["결과"]).name)
        except Exception as e:
            failed.append({"파일": src.name, "이유": str(e)})

    _log(f"[file-convert] 영상 {len(done)}/{len(files)}개 변환 → {out_root}")
    summary = f"{_human_size(total_before)} → {_human_size(total_after)}"
    if total_before:
        change = (1 - total_after / total_before) * 100
        summary += f" ({change:.0f}% 절감)" if change >= 0 else f" ({-change:.0f}% 증가)"
    return {
        "변환 완료": f"{len(done)}개",
        "실패": failed,
        "저장 위치": str(out_root),
        "전체 용량": summary,
        "파일 목록": done[:50],
    }


# ─────────────────────────────────────────────
# 3단계: 문서 변환
#  - PDF → DOCX / 이미지, 이미지 → PDF, XLSX ↔ CSV : pip 패키지만으로 동작
#  - HWP·HWPX·DOCX·PPTX·XLSX → PDF, HWP → DOCX : LibreOffice 필요 (무료, 한 번만 설치)
# ─────────────────────────────────────────────

OFFICE_INPUTS = {
    ".hwp", ".hwpx", ".doc", ".docx", ".odt", ".rtf", ".txt",
    ".ppt", ".pptx", ".odp", ".key",
    ".xls", ".xlsx", ".ods",
}
HWP_INPUTS = {".hwp", ".hwpx"}
WRITER_INPUTS = {".hwp", ".hwpx", ".doc", ".docx", ".odt", ".rtf", ".txt"}
SHEET_INPUTS = {".xls", ".xlsx", ".ods"}
DOC_INPUT_EXTENSIONS = OFFICE_INPUTS | {".pdf", ".csv"}

DOC_ROUTES_HELP = [
    "hwp · hwpx · doc · docx · pptx · xlsx 등 → pdf",
    "hwp · hwpx · doc · odt · rtf → docx",
    "pdf → docx",
    "pdf → png · jpg (페이지별 이미지)",
    "xlsx → csv (시트별), csv → xlsx",
    "이미지 여러 장 → pdf 하나 (images_to_pdf)",
]

# HWP를 읽게 해주는 LibreOffice 확장 기능 (H2Orestart, GPL-3.0)
H2O_URL = "https://raw.githubusercontent.com/ebandal/H2Orestart/main/dist/H2Orestart.oxt"
APP_DIR = Path.home() / ".file-converter-mcp"


def _find_soffice() -> str | None:
    candidates = [
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        str(Path.home() / "Applications/LibreOffice.app/Contents/MacOS/soffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
    ]
    for c in candidates:
        if Path(c).exists():
            return c
    return shutil.which("soffice") or shutil.which("libreoffice")


def _require_soffice() -> str:
    soffice = _find_soffice()
    if soffice:
        return soffice
    if platform.system() == "Darwin":
        how = (
            "터미널에서 `brew install --cask libreoffice` 를 실행하거나, "
            "https://www.libreoffice.org/download 에서 macOS용을 받아 설치해주세요. "
        )
    else:
        how = "https://www.libreoffice.org/download 에서 LibreOffice를 설치해주세요. "
    raise ToolError(
        "HWP·오피스 문서를 변환하려면 무료 프로그램 LibreOffice가 한 번 설치돼 있어야 해요. "
        + how + "설치 후 다시 요청하면 됩니다."
    )


def _unopkg_path(soffice: str) -> str | None:
    p = Path(soffice).with_name("unopkg.exe" if soffice.endswith(".exe") else "unopkg")
    return str(p) if p.exists() else shutil.which("unopkg")


def _ensure_hwp_support(soffice: str) -> None:
    """HWP 확장 기능이 없으면 처음 한 번 자동으로 내려받아 설치"""
    unopkg = _unopkg_path(soffice)
    if not unopkg:
        raise ToolError("LibreOffice의 확장 관리 도구(unopkg)를 찾지 못했어요. LibreOffice를 다시 설치해주세요.")

    listed = subprocess.run([unopkg, "list"], capture_output=True, text=True, timeout=120)
    shared = subprocess.run([unopkg, "list", "--shared"], capture_output=True, text=True, timeout=120)
    if "H2Orestart" in (listed.stdout + shared.stdout):
        return

    APP_DIR.mkdir(parents=True, exist_ok=True)
    oxt = APP_DIR / "H2Orestart.oxt"
    if not oxt.exists():
        _log("[file-convert] HWP 확장 기능 내려받는 중...")
        try:
            urllib.request.urlretrieve(H2O_URL, oxt)
        except Exception as e:
            oxt.unlink(missing_ok=True)
            raise ToolError(f"HWP 확장 기능을 내려받지 못했어요. 인터넷 연결을 확인해주세요. ({e})")

    _log("[file-convert] HWP 확장 기능 설치 중...")
    cmd = [unopkg, "add", "--suppress-license", str(oxt)]
    if hasattr(os, "geteuid") and os.geteuid() == 0:
        cmd.insert(2, "--shared")  # 관리자(root) 계정에서는 공용 설치만 허용됨
    proc = subprocess.run(
        cmd,
        capture_output=True, text=True, timeout=300,
    )
    if proc.returncode != 0:
        raise ToolError(
            "HWP 확장 기능 설치에 실패했어요. LibreOffice 창이 열려 있다면 닫고 다시 시도해주세요. "
            f"({(proc.stderr or proc.stdout).strip()[-200:]})"
        )


def _soffice_convert(src: Path, out_dir: Path, target: str) -> Path:
    """LibreOffice로 변환. target 예: 'pdf', 'docx:MS Word 2007 XML'"""
    soffice = _require_soffice()
    if src.suffix.lower() in HWP_INPUTS:
        _ensure_hwp_support(soffice)

    ext = target.split(":")[0]
    out_dir.mkdir(parents=True, exist_ok=True)
    # 같은 이름 덮어쓰기 방지를 위해 임시 폴더에 만든 뒤 옮김
    with tempfile.TemporaryDirectory() as tmp:
        try:
            proc = subprocess.run(
                [soffice, "--headless", "--norestore", "--convert-to", target, "--outdir", tmp, str(src)],
                capture_output=True, text=True, timeout=300,
            )
        except subprocess.TimeoutExpired:
            raise ToolError(f"{src.name} 변환 시간이 너무 오래 걸려 중단했어요.")
        produced = Path(tmp) / f"{src.stem}.{ext}"
        if not produced.exists():
            detail = (proc.stderr or proc.stdout or "").strip().splitlines()
            raise ToolError(
                f"{src.name} 변환 실패. LibreOffice 창이 열려 있다면 닫고 다시 시도해주세요."
                + (f" ({detail[-1][:200]})" if detail else "")
            )
        dst = _unique_path(out_dir / produced.name)
        shutil.move(str(produced), dst)
    return dst


def _pdf_to_docx(src: Path, out_dir: Path) -> Path:
    from pdf2docx import Converter

    out_dir.mkdir(parents=True, exist_ok=True)
    dst = _unique_path(out_dir / f"{src.stem}.docx")
    cv = Converter(str(src))
    try:
        cv.convert(str(dst))
    finally:
        cv.close()
    return dst


def _pdf_to_images(src: Path, out_dir: Path, fmt: str, dpi: int) -> list[Path]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(src))
    target = _unique_path(out_dir / src.stem)
    target.mkdir(parents=True, exist_ok=True)
    results = []
    try:
        digits = len(str(len(pdf)))
        for i, page in enumerate(pdf, start=1):
            img = page.render(scale=dpi / 72).to_pil()
            if fmt == "jpg":
                img = img.convert("RGB")
            path = target / f"{src.stem}_{str(i).zfill(digits)}.{fmt}"
            img.save(path, "JPEG" if fmt == "jpg" else "PNG", **({"quality": 92} if fmt == "jpg" else {}))
            results.append(path)
    finally:
        pdf.close()
    return results


def _xlsx_to_csv(src: Path, out_dir: Path) -> list[Path]:
    from openpyxl import load_workbook

    wb = load_workbook(src, read_only=True, data_only=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    results = []
    try:
        multi = len(wb.sheetnames) > 1
        for ws in wb.worksheets:
            name = f"{src.stem}_{ws.title}.csv" if multi else f"{src.stem}.csv"
            dst = _unique_path(out_dir / name)
            # utf-8-sig: 윈도우 엑셀에서 열어도 한글이 안 깨짐
            with open(dst, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                for row in ws.iter_rows(values_only=True):
                    writer.writerow(["" if v is None else v for v in row])
            results.append(dst)
    finally:
        wb.close()
    return results


def _read_csv_rows(src: Path) -> list[list[str]]:
    for enc in ("utf-8-sig", "cp949", "euc-kr", "utf-8"):
        try:
            with open(src, newline="", encoding=enc) as f:
                return list(csv.reader(f))
        except UnicodeDecodeError:
            continue
    raise ToolError(f"{src.name}의 글자 인코딩을 읽지 못했어요.")


def _cell_value(v: str):
    """숫자는 숫자로, 나머지는 글자로. 0으로 시작하는 값(전화번호·우편번호)은 글자로 유지"""
    t = v.strip()
    if not t or (len(t) > 1 and t.startswith("0") and not t.startswith("0.")):
        return v
    try:
        return int(t)
    except ValueError:
        pass
    try:
        return float(t)
    except ValueError:
        return v


def _csv_to_xlsx(src: Path, out_dir: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = src.stem[:31] or "Sheet1"
    for row in _read_csv_rows(src):
        ws.append([_cell_value(v) for v in row])
    out_dir.mkdir(parents=True, exist_ok=True)
    dst = _unique_path(out_dir / f"{src.stem}.xlsx")
    wb.save(dst)
    return dst


def _convert_document_one(src: Path, fmt: str, out_dir: Path, dpi: int) -> dict:
    ext = src.suffix.lower()
    outputs: list[Path]

    if ext == ".pdf":
        if fmt == "docx":
            outputs = [_pdf_to_docx(src, out_dir)]
        elif fmt in ("png", "jpg"):
            outputs = _pdf_to_images(src, out_dir, fmt, dpi)
        else:
            raise ToolError("PDF는 docx, png, jpg로 바꿀 수 있어요.")
    elif ext == ".csv":
        if fmt != "xlsx":
            raise ToolError("CSV는 xlsx로 바꿀 수 있어요.")
        outputs = [_csv_to_xlsx(src, out_dir)]
    elif ext in OFFICE_INPUTS:
        if fmt == "pdf":
            outputs = [_soffice_convert(src, out_dir, "pdf")]
        elif fmt == "docx" and ext in WRITER_INPUTS and ext != ".docx":
            outputs = [_soffice_convert(src, out_dir, "docx:MS Word 2007 XML")]
        elif fmt == "csv" and ext == ".xlsx":
            outputs = _xlsx_to_csv(src, out_dir)
        elif fmt == "csv" and ext in SHEET_INPUTS:
            outputs = [_soffice_convert(src, out_dir, "csv:Text - txt - csv (StarCalc):44,34,76,1")]
        elif fmt == "xlsx" and ext in SHEET_INPUTS and ext != ".xlsx":
            outputs = [_soffice_convert(src, out_dir, "xlsx:Calc MS Excel 2007 XML")]
        elif fmt == "pptx" and ext in {".ppt", ".odp", ".key"}:
            outputs = [_soffice_convert(src, out_dir, "pptx:Impress MS PowerPoint 2007 XML")]
        else:
            raise ToolError(f"{ext.lstrip('.')} → {fmt} 변환은 아직 지원하지 않아요. 가능한 경로: " + " / ".join(DOC_ROUTES_HELP))
    else:
        raise ToolError(f"문서 파일이 아니거나 아직 지원하지 않는 형식이에요: {src.name}")

    before = src.stat().st_size
    after = sum(p.stat().st_size for p in outputs)
    result = {
        "원본": str(src),
        "결과": str(outputs[0]) if len(outputs) == 1 else f"{len(outputs)}개 파일 → {outputs[0].parent}",
        "용량": f"{_human_size(before)} → {_human_size(after)}",
        "_before": before,
        "_after": after,
        "_names": [p.name for p in outputs],
    }
    if src.suffix.lower() in HWP_INPUTS:
        result["참고"] = "HWP는 표·글꼴·줄 간격이 원본과 조금 다를 수 있어요. 중요한 문서는 결과를 한 번 확인해주세요."
    return result


@mcp.tool()
def convert_document(
    path: str,
    to_format: str,
    dpi: int = 150,
    output_dir: str | None = None,
) -> dict:
    """
    문서 파일 하나를 변환합니다. 원본은 그대로 둡니다.
    - hwp · hwpx · doc · docx · pptx · xlsx 등 → pdf
    - hwp · hwpx · doc · odt · rtf → docx (워드에서 편집 가능)
    - pdf → docx, pdf → png / jpg (페이지별 이미지)
    - xlsx → csv (시트가 여러 개면 시트별로), csv → xlsx
    HWP·오피스 → PDF/DOCX는 LibreOffice가 필요하며, HWP 확장 기능은 처음 한 번 자동 설치됩니다.

    Args:
        path: 변환할 파일 경로 (예: ~/Downloads/공고문.hwp)
        to_format: pdf, docx, png, jpg, csv, xlsx, pptx
        dpi: PDF → 이미지 변환 시 해상도 (기본 150, 인쇄용은 300)
        output_dir: 저장할 폴더. 비우면 원본 옆 'converted_형식' 폴더
    """
    src = Path(path).expanduser().resolve()
    if not src.is_file():
        raise ToolError(f"파일을 찾을 수 없어요: {src}")
    fmt = to_format.lower().lstrip(".").strip()
    fmt = "jpg" if fmt == "jpeg" else fmt
    out_dir = Path(output_dir).expanduser().resolve() if output_dir else src.parent / f"converted_{fmt}"

    result = _convert_document_one(src, fmt, out_dir, dpi)
    _log(f"[file-convert] {src.name} → {result['결과']}")
    return {k: v for k, v in result.items() if not k.startswith("_")}


@mcp.tool()
def convert_document_folder(
    folder: str,
    to_format: str,
    from_format: str | None = None,
    include_subfolders: bool = False,
    output_dir: str | None = None,
) -> dict:
    """
    폴더 안의 문서를 한 번에 변환합니다. 원본은 그대로 둡니다.
    예: 폴더 안 hwp 전부 pdf로, docx 전부 pdf로, xlsx 전부 csv로

    Args:
        folder: 문서가 들어 있는 폴더 경로
        to_format: pdf, docx, png, jpg, csv, xlsx
        from_format: 특정 형식만 변환할 때 지정 (예: hwp). 비우면 바꿀 수 있는 문서 전부
        include_subfolders: 하위 폴더까지 포함할지 여부
        output_dir: 저장할 폴더. 비우면 대상 폴더 안 'converted_형식' 폴더
    """
    base = Path(folder).expanduser().resolve()
    if not base.is_dir():
        raise ToolError(f"폴더를 찾을 수 없어요: {base}")
    fmt = to_format.lower().lstrip(".").strip()
    fmt = "jpg" if fmt == "jpeg" else fmt
    only = ("." + from_format.lower().lstrip(".")) if from_format else None
    out_root = Path(output_dir).expanduser().resolve() if output_dir else base / f"converted_{fmt}"

    pattern = "**/*" if include_subfolders else "*"
    files = sorted(
        p for p in base.glob(pattern)
        if p.is_file()
        and p.suffix.lower() in DOC_INPUT_EXTENSIONS
        and (only is None or p.suffix.lower() == only)
        and p.suffix.lower() != f".{fmt}"
        and not p.name.startswith("~$")  # 열려 있는 오피스 문서의 임시 파일 제외
        and out_root not in p.parents
        and not any(part.startswith("converted_") for part in p.relative_to(base).parts[:-1])
    )
    if not files:
        return {"메시지": "변환할 문서가 없어요.", "폴더": str(base)}

    done, failed = [], []
    total_before = total_after = 0
    for src in files:
        target_dir = out_root / src.parent.relative_to(base)
        try:
            r = _convert_document_one(src, fmt, target_dir, 150)
            total_before += r["_before"]
            total_after += r["_after"]
            done.extend(r["_names"])
        except Exception as e:
            failed.append({"파일": src.name, "이유": str(e)})

    _log(f"[file-convert] 문서 {len(files) - len(failed)}/{len(files)}개 변환 → {out_root}")
    return {
        "변환 완료": f"{len(files) - len(failed)}개",
        "실패": failed,
        "저장 위치": str(out_root),
        "전체 용량": f"{_human_size(total_before)} → {_human_size(total_after)}",
        "파일 목록": done[:50],
    }


@mcp.tool()
def images_to_pdf(
    paths: list[str] | None = None,
    folder: str | None = None,
    output_name: str | None = None,
    output_dir: str | None = None,
) -> dict:
    """
    이미지 여러 장을 PDF 파일 하나로 묶습니다. (스캔본·사진 제출용)
    paths로 파일 목록을 주거나, folder로 폴더 안 이미지 전체(파일명 순)를 지정하세요.

    Args:
        paths: 이미지 파일 경로 목록 (이 순서대로 페이지가 됨)
        folder: 이미지가 들어 있는 폴더 경로 (파일명 순서로 묶음)
        output_name: 만들 PDF 이름 (확장자 없이). 비우면 폴더명 또는 첫 파일명
        output_dir: 저장할 폴더. 비우면 이미지가 있는 곳의 'converted_pdf' 폴더
    """
    if paths:
        files = [Path(p).expanduser().resolve() for p in paths]
        missing = [str(p) for p in files if not p.is_file()]
        if missing:
            raise ToolError(f"파일을 찾을 수 없어요: {', '.join(missing[:5])}")
        base = files[0].parent
        default_name = files[0].stem
    elif folder:
        base = Path(folder).expanduser().resolve()
        if not base.is_dir():
            raise ToolError(f"폴더를 찾을 수 없어요: {base}")
        files = sorted(p for p in base.iterdir() if p.is_file() and p.suffix.lower() in INPUT_EXTENSIONS)
        default_name = base.name
    else:
        raise ToolError("paths(파일 목록)나 folder(폴더) 중 하나를 지정해주세요.")
    if not files:
        raise ToolError("PDF로 묶을 이미지가 없어요.")

    pages = []
    for p in files:
        with Image.open(p) as img:
            img = ImageOps.exif_transpose(img)
            if img.mode in ("RGBA", "LA", "P"):
                img = img.convert("RGBA")
                bg = Image.new("RGB", img.size, (255, 255, 255))
                bg.paste(img, mask=img.split()[-1])
                img = bg
            else:
                img = img.convert("RGB")
            pages.append(img.copy())

    out_dir = Path(output_dir).expanduser().resolve() if output_dir else base / "converted_pdf"
    out_dir.mkdir(parents=True, exist_ok=True)
    dst = _unique_path(out_dir / f"{output_name or default_name}.pdf")
    pages[0].save(dst, "PDF", save_all=True, append_images=pages[1:], resolution=150)
    _log(f"[file-convert] 이미지 {len(pages)}장 → {dst}")
    return {"결과": str(dst), "페이지 수": len(pages), "용량": _human_size(dst.stat().st_size)}


if __name__ == "__main__":
    mcp.run()
